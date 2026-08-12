"""
memo.py — the financing memorandum.

Produces the same memo in two formats from one set of prose:

    memo.md    Markdown, for review and version control
    memo.docx  Word, for circulation

The memo follows the structure Tierra uses on Colorado metro district
financings — transaction summary, security and pledged revenue, development
and absorption, taxable value build, sizing results, sensitivities, and the
assumption appendix — with the Utah statutory framework substituted for
Colorado's.
"""

from __future__ import annotations

import os
from dataclasses import dataclass
from datetime import date
from typing import Iterable, Optional, Sequence

from .config import PID_STATUTORY_LEVY_CAP
from .engine import Model, Results

# ── Small formatting helpers ─────────────────────────────────────────────────

def _m(v: Optional[float]) -> str:
    return "—" if v is None else f"${v:,.0f}"


def _mm(v: Optional[float]) -> str:
    return "—" if v is None else f"${v / 1e6:,.2f}MM"


def _p(v: Optional[float], places: int = 2) -> str:
    return "—" if v is None else f"{v:.{places}%}"


def _x(v: Optional[float]) -> str:
    return "—" if v is None else f"{v:.2f}x"


def _d(v: Optional[date]) -> str:
    return "—" if v is None else v.strftime("%B %-d, %Y")


@dataclass
class Table:
    title: str
    headers: Sequence[str]
    rows: Sequence[Sequence[str]]
    note: str = ""


@dataclass
class Section:
    heading: str
    paragraphs: Sequence[str] = ()
    bullets: Sequence[str] = ()
    tables: Sequence[Table] = ()


# ── Content ───────────────────────────────────────────────────────────────────

def _sensitivities(res: Results) -> Table:
    """Re-run the model under the standard downside cases."""
    from dataclasses import replace

    cases = [
        ("Base case", {}),
        ("Absorption −25%", {"hypothetical_scenario": "Yes", "absorption_scenario": 0.75,
                             "lot_delivery_scenario": 0.75}),
        ("No home price inflation", {"inflation_rate": 0.0}),
        ("No reassessment growth", {"reassess_rate_resid": 0.0,
                                    "reassess_rate_resid_sub": 0.0}),
        ("Collections at 95%", {"tax_collect_mill_prc": 0.95}),
        ("Coverage at 1.50x", {"dsc_senior_lien_bonds": 1.50}),
    ]
    rows = []
    for label, overrides in cases:
        try:
            cfg = replace(res.cfg, **overrides) if overrides else res.cfg
            out = res if not overrides else Model(cfg, res.dev).run()
        except Exception:                      # a case that cannot be sized
            rows.append([label, "n/a", "n/a", "n/a", "n/a"])
            continue
        rows.append([
            label, _m(out.senior_par), _m(out.sub_par),
            _m(out.total_reimbursement), _m(out.reimbursement_per_lot),
        ])
    return Table(
        title="Sensitivity analysis",
        headers=["Case", "Senior par", "Subordinate par",
                 "Total reimbursement", "Per lot"],
        rows=rows,
        note="Each case re-runs the full model; only the named assumption changes.",
    )


def build_sections(res: Results) -> list[Section]:
    cfg, dev = res.cfg, res.dev
    s, sub = res.senior_stats, res.sub_stats
    statutory_mills = PID_STATUTORY_LEVY_CAP * 1000

    # Milestones from the projection.
    first_rev = next((r for r in res.summary if r.senior_levy_collections > 0), None)
    buildout = next((r for r in reversed(res.summary) if r.residential_units > 0), None)
    stabilised = max(res.summary, key=lambda r: r.senior_taxable_value)

    sections: list[Section] = []

    sections.append(Section(
        heading="1. Transaction summary",
        paragraphs=[
            f"{cfg.district_name} (the “District”) is a public infrastructure "
            f"district organized under the Utah Public Infrastructure District Act, "
            f"Title 17D, Chapter 4, Utah Code, within {cfg.city} City, "
            f"{cfg.county} County, Utah. This memorandum sets out the projected "
            f"capacity of the District's limited tax general obligation bonds at a "
            f"debt service levy of {cfg.mill_levy_ds_target:.3f} mills, and the "
            f"reimbursement the financing is expected to deliver to "
            f"{cfg.developer} (the “Developer”).",

            f"On a delivery date of {_d(cfg.delivery)}, the model sizes "
            f"{_mm(res.senior_par)} of {cfg.senior_bonds_series} senior lien bonds "
            f"at {_p(cfg.senior_interest_rate, 3)} and {_mm(res.sub_par)} of "
            f"{cfg.sub_bonds_series} subordinate lien cashflow bonds at "
            f"{_p(cfg.sub_interest_rate, 3)}. Together they produce "
            f"{_m(res.total_reimbursement)} of reimbursement to the Developer — "
            f"approximately {_m(res.reimbursement_per_lot)} per residential unit "
            f"across {dev.total_units():,} units.",
        ],
        tables=[Table(
            title="Sources and uses",
            headers=["", "Senior lien", "Subordinate lien", "Total"],
            rows=[
                ["Par amount", _m(res.senior_par), _m(res.sub_par), _m(res.total_par)],
                ["Premium / (discount)", _m(s.premium), _m(0), _m(s.premium)],
                ["Total sources", _m(res.senior_par + s.premium), _m(res.sub_par),
                 _m(res.total_par + s.premium)],
                ["", "", "", ""],
                ["Developer reimbursement", _m(res.senior_reimbursement),
                 _m(res.sub_reimbursement), _m(res.total_reimbursement)],
                ["Debt service reserve / surplus fund", _m(res.surplus_fund_deposit),
                 _m(0), _m(res.surplus_fund_deposit)],
                ["Capitalized interest", _m(res.capitalized_interest_deposit),
                 _m(0), _m(res.capitalized_interest_deposit)],
                ["Underwriters' discount", _m(res.senior_uwd), _m(res.sub_uwd),
                 _m(res.senior_uwd + res.sub_uwd)],
                ["Costs of issuance", _m(cfg.coi), _m(0), _m(cfg.coi)],
                ["Total uses", _m(res.senior_par + s.premium), _m(res.sub_par),
                 _m(res.total_par + s.premium)],
            ],
        )],
    ))

    sections.append(Section(
        heading="2. Security and pledged revenue",
        paragraphs=[
            f"The bonds are limited tax general obligations of the District, "
            f"payable from an ad valorem property tax levied against all taxable "
            f"property within the District's boundaries. The levy is capped at "
            f"three levels, the most restrictive of which controls: "
            f"{statutory_mills:,.3f} mills (0.015 per dollar of taxable value) "
            f"under Section 17D-4-303, Utah Code; the rate fixed in the District's "
            f"governing document; and the rate fixed in the indentures. This "
            f"analysis assumes the controlling rate is "
            f"{cfg.mill_levy_governing_doc:.3f} mills.",

            "Because the levy securing the required mill levy does not exceed the "
            "rate established in the Public Infrastructure District Act, the "
            "governing document, or the indentures, the District is not required "
            "to give Truth in Taxation notice or hold a public hearing to impose "
            "it. That distinguishes a Utah PID levy from a levy above the "
            "certified tax rate, and removes the annual political risk a Colorado "
            "metropolitan district faces when a service plan cap is approached.",

            f"Property taxes in Utah are levied against taxable value determined as "
            f"of 1 January and are due on 30 November of the same year. Debt "
            f"service is therefore structured with principal due "
            f"{date(2000, cfg.prin_maturity, cfg.prin_maturity_day_senior).strftime('%-d %B')} "
            f"and interest semi-annually, so each year's collections are in hand "
            f"before the payment they support. Delinquent taxes carry a penalty of "
            f"the greater of 2.5% or $10, and the Act permits the District to "
            f"impose an additional 7% annual penalty. This analysis assumes "
            f"collections of {_p(cfg.tax_collect_mill_prc)} of the levy and takes "
            f"no credit for penalty revenue.",
        ],
        bullets=[
            f"Senior lien coverage requirement: {_x(cfg.dsc_senior_lien_bonds)} of "
            f"net pledged revenue.",
            f"Subordinate lien coverage requirement: {_x(cfg.dsc_sub_lien_bonds)}; "
            f"the subordinate bonds are cashflow bonds, paid only from revenue "
            f"released after senior debt service, with unpaid interest accruing at "
            f"{_p(cfg.sub_interest_rate, 3)}.",
            f"Debt service reserve / surplus fund: {_m(res.surplus_fund_deposit)}, "
            f"sized at the least of 10% of par, 125% of average annual debt "
            f"service, and maximum annual debt service.",
            f"Capitalized interest: {_m(res.capitalized_interest_deposit)}, funding "
            f"interest through {_d(cfg.capi_end_date)} "
            f"({cfg.capi_term_months} months from closing).",
            f"Personal property uniform fees (Section 59-2-405, Utah Code) are "
            f"distributed to the District in the same proportion as real property "
            f"tax; this analysis credits {_p(cfg.uniform_fee_prc)} of the levy.",
        ],
    ))

    absorb_rows = [
        [p.name, p.product_type, f"{p.total_units:,}", _m(p.asp),
         f"{min(p.home_closings) if p.home_closings else '—'}",
         f"{max(p.home_closings) if p.home_closings else '—'}",
         f"{p.average_absorption():,.0f}"]
        for p in dev.products if p.total_units
    ]
    absorb_rows.append(["Total", "", f"{dev.total_units():,}", "", "", "", ""])

    sections.append(Section(
        heading="3. Development program and absorption",
        paragraphs=[
            f"The development comprises {dev.total_units():,} residential units "
            f"across {len([p for p in dev.products if p.total_units])} product "
            f"types. Finished lots are delivered to the builder roughly "
            f"{cfg.home_lot_delivery_lead_months} months before the first home "
            f"closing, and closings run from "
            f"{min((min(p.home_closings) for p in dev.products if p.home_closings), default='—')} "
            f"through "
            f"{max((max(p.home_closings) for p in dev.products if p.home_closings), default='—')}. "
            f"Home prices are inflated at {_p(cfg.inflation_rate)} per year from "
            f"the {cfg.resid_delivery_year.year} base.",
        ],
        tables=[Table(
            title="Residential program",
            headers=["Product", "Type", "Units", "Base ASP",
                     "First closing", "Last closing", "Avg. annual absorption"],
            rows=absorb_rows,
        )],
    ))

    tv_rows = [
        [str(r.assessment_date.year), f"{r.lot_units:,.0f}",
         f"{r.residential_units:,.0f}", _m(r.lot_taxable_value),
         _m(r.new_home_taxable_value), _m(r.senior_taxable_value),
         _m(r.senior_net_revenue)]
        for r in res.summary[:20]
    ]

    sections.append(Section(
        heading="4. Taxable value",
        paragraphs=[
            "Utah taxes primary residential property on 55% of fair market value — "
            "the 45% primary residential exemption under Section 59-2-103, Utah "
            "Code, which reaches up to one acre of land per residential unit. "
            "Finished lots held in builder inventory are carried at the same "
            "ratio: Utah Admin. Code R884-24P-52 allows the residential exemption "
            "on unoccupied property the county assessor determines will qualify as "
            "a primary residence once occupied. This is the single largest "
            "difference from the Colorado template, where homes are assessed at "
            "the residential rate (6.7% in the reference model) and vacant land at "
            "the 29% non-residential rate.",

            f"County assessors revalue annually, so value created in one calendar "
            f"year appears on the following year's roll — a "
            f"{cfg.value_lag_years}-year lag, against the two-year lag Colorado's "
            f"biennial reassessment cycle produces. Values already on the roll are "
            f"grown at {_p(cfg.reassess_rate_resid)} per year. Finished lots are "
            f"carried at {_p(cfg.platted_lot_value)} of the eventual home price "
            f"until a home closes on them.",

            f"Taxable value first supports a levy in "
            f"{first_rev.assessment_date.year if first_rev else '—'}, reaches "
            f"{_m(buildout.senior_taxable_value) if buildout else '—'} at buildout "
            f"in {buildout.assessment_date.year if buildout else '—'}, and grows "
            f"to {_m(stabilised.senior_taxable_value)} by "
            f"{stabilised.assessment_date.year} on reassessment alone.",
        ],
        tables=[Table(
            title="Taxable value and pledged revenue build",
            headers=["Assessment year", "Lots delivered", "Homes closed",
                     "Lot taxable value", "Home taxable value",
                     "Total taxable value", "Net pledged revenue"],
            rows=tv_rows,
            note="First twenty projection years; the full series is on the Summary tab.",
        )],
    ))

    ds_rows = [
        [str(r.payment_date.year), _m(r.principal), _m(r.interest * 2),
         _m(r.annual_net), _m(r.revenue), _x(r.actual_coverage)]
        for r in res.senior
        if r.payment_date.month == cfg.prin_maturity and (r.annual_net or r.principal)
    ]

    sections.append(Section(
        heading="5. Sizing results and debt service",
        paragraphs=[
            f"Principal is solved backwards from the final maturity: in each year "
            f"the projected net pledged revenue is divided by the "
            f"{_x(cfg.dsc_senior_lien_bonds)} coverage requirement, the coupon "
            f"generated by later maturities is removed, and the balance is turned "
            f"into principal in $5,000 denominations. The result is "
            f"{_mm(res.senior_par)} of senior bonds with an average life of "
            f"{s.average_life:.2f} years, final maturity {_d(s.final_maturity)}, "
            f"maximum annual debt service of {_m(s.max_annual_debt_service)} and "
            f"total debt service of {_m(s.total_debt_service)}.",

            f"The arbitrage TIC is {_p(s.arbitrage_tic, 3)} and the all-in TIC, "
            f"including the underwriters' discount and costs of issuance, is "
            f"{_p(s.all_in_tic, 3)}. Subordinate bonds of {_mm(res.sub_par)} are "
            f"sized to the largest amount the residual cashflow retires in full — "
            f"principal and accrued interest — by {_d(sub.final_maturity)}. Total "
            f"debt service across both liens is "
            f"{_m(s.total_debt_service + sub.total_debt_service)}, a repayment "
            f"ratio of {_x(res.repayment_ratio)} on the reimbursement delivered.",
        ],
        tables=[Table(
            title="Senior lien debt service",
            headers=["Year", "Principal", "Interest", "Net debt service",
                     "Net pledged revenue", "Coverage"],
            rows=ds_rows,
            note="Net of capitalized interest and surplus fund earnings.",
        )],
    ))

    sections.append(Section(
        heading="6. Sensitivities",
        paragraphs=[
            "The capacity of the financing is driven by absorption pace, home "
            "prices and the reassessment assumption. The cases below re-run the "
            "full model with one assumption changed at a time.",
        ],
        tables=[_sensitivities(res)],
    ))

    sections.append(Section(
        heading="7. How this differs from the Colorado model",
        paragraphs=[
            "The model is the Colorado metropolitan district template — same tabs, "
            "same layout, same sizing mechanics. The substantive changes are the "
            "ones Utah law requires.",
        ],
        tables=[Table(
            title="Colorado metropolitan district vs. Utah PID",
            headers=["Item", "Colorado", "Utah (this model)"],
            rows=[
                ["Enabling act", "Title 32, Article 1, C.R.S. (special districts)",
                 "Title 17D, Chapter 4, Utah Code (PID Act)"],
                ["Levy cap", "Service plan mill levy cap, adjusted for changes in "
                             "the residential assessment rate",
                 f"Least of {statutory_mills:,.3f} mills by statute, the governing "
                 f"document cap, and the indenture cap — no adjustment mechanism"],
                ["Residential assessment",
                 "Gallagher / statutory residential assessment rate (6.7% in the "
                 "reference model); vacant land at 29%",
                 f"45% primary residential exemption → {_p(cfg.resid_taxable_ratio)} "
                 f"of fair market value; builder inventory at the same ratio under "
                 f"R884-24P-52"],
                ["Levy cap adjustment", "Gallagherization of the service plan cap",
                 "None — the cap is a fixed rate per dollar of taxable value"],
                ["Reassessment cycle", "Biennial",
                 f"Annual → a {cfg.value_lag_years}-year lag from value creation "
                 f"to the tax roll"],
                ["Tax due dates", "Half 28 February / half 15 June (or full 30 April)",
                 "Single payment, 30 November"],
                ["Principal payment date", "1 December",
                 f"{date(2000, cfg.prin_maturity, cfg.prin_maturity_day_senior).strftime('%-d %B')}"],
                ["Vehicle tax revenue",
                 "Specific ownership tax, ~6–8% of the levy",
                 f"Personal property uniform fee (Section 59-2-405), distributed "
                 f"pro rata; credited at {_p(cfg.uniform_fee_prc)}"],
                ["County collection cost",
                 "County treasurer fee deducted from the distribution (~1.5%)",
                 f"Recovered through a separate statewide levy on property "
                 f"(Section 59-2-1602); {_p(cfg.county_treasurer_fee)} deducted here"],
                ["Rate increase procedure",
                 "TABOR election for new mill levies",
                 "Truth in Taxation hearing above the certified tax rate; the "
                 "required mill levy is exempt while within the caps"],
                ["Operations levy",
                 "Separate operations and maintenance mill levy",
                 "None assumed — district administration is charged against "
                 "pledged revenue"],
                ["Agricultural land", "Agricultural classification",
                 "Greenbelt Reduction (Section 59-2-503) with up to five years of "
                 "rollback tax on withdrawal — a Developer obligation"],
            ],
        )],
    ))

    appendix_rows = [
        ["Delivery date", _d(cfg.delivery)],
        ["First interest — senior / subordinate",
         f"{_d(cfg.first_int)} / {_d(cfg.first_int_sub)}"],
        ["First principal", _d(s.first_maturity)],
        ["Final maturity — senior / subordinate",
         f"{_d(s.final_maturity)} / {_d(sub.final_maturity)}"],
        ["Premium call / par call",
         f"{_d(cfg.premium_call_first)} at {cfg.premium_call_price:.0f} / "
         f"{_d(cfg.par_call_first)}"],
        ["Senior / subordinate interest rate",
         f"{_p(cfg.senior_interest_rate, 3)} / {_p(cfg.sub_interest_rate, 3)}"],
        ["Senior / subordinate coverage",
         f"{_x(cfg.dsc_senior_lien_bonds)} / {_x(cfg.dsc_sub_lien_bonds)}"],
        ["Debt service mill levy", f"{cfg.mill_levy_ds_target:.3f} mills"],
        ["Statutory levy cap (17D-4-303)", f"{statutory_mills:,.3f} mills"],
        ["Primary residential taxable ratio", _p(cfg.resid_taxable_ratio)],
        ["Developed lot taxable ratio", _p(cfg.developed_lot_value)],
        ["Finished lot value (% of ASP)", _p(cfg.platted_lot_value)],
        ["Property tax collection rate", _p(cfg.tax_collect_mill_prc)],
        ["Annual reassessment — existing / new",
         f"{_p(cfg.reassess_rate_resid)} / {_p(cfg.reassess_rate_resid_sub)}"],
        ["Home price inflation", _p(cfg.inflation_rate)],
        ["Value lag", f"{cfg.value_lag_years} year(s)"],
        ["Underwriters' discount — senior / subordinate",
         f"{_p(cfg.uwd_senior)} / {_p(cfg.uwd_sub)}"],
        ["Costs of issuance", _m(cfg.coi)],
        ["Trustee fee — senior / subordinate",
         f"{_m(cfg.trustee_fee)} / {_m(cfg.trustee_fee_sub)}"],
        ["Annual district administration",
         f"{_m(cfg.admin_cost_base)} growing at {_p(cfg.admin_cost_growth)}"],
        ["Interest earnings rate", _p(cfg.interest_earn_rate)],
        ["Capitalized interest period", f"{cfg.capi_term_months} months"],
    ]

    sections.append(Section(
        heading="Appendix A — Key assumptions",
        tables=[Table(title="", headers=["Assumption", "Value"], rows=appendix_rows)],
    ))

    if res.warnings:
        sections.append(Section(
            heading="Appendix B — Model notes",
            bullets=list(res.warnings),
        ))

    return sections


# ── Renderers ─────────────────────────────────────────────────────────────────

def _memo_header(res: Results) -> tuple[str, list[tuple[str, str]]]:
    cfg = res.cfg
    title = "Financing Memorandum"
    meta = [
        ("To", cfg.developer),
        ("From", "Tierra Financial Advisors, LLC"),
        ("Re", f"{cfg.district_name} — Limited Tax General Obligation Bonds"),
        ("Levy", f"{cfg.mill_levy_ds_target:.3f} mills for debt service"),
        ("Date", date.today().strftime("%B %-d, %Y")),
    ]
    return title, meta


def write_markdown(res: Results, path: str) -> str:
    title, meta = _memo_header(res)
    out: list[str] = [f"# {title}", ""]
    out += [f"**{k}:** {v}  " for k, v in meta]
    out += ["", "---", ""]

    for section in build_sections(res):
        out += [f"## {section.heading}", ""]
        for para in section.paragraphs:
            out += [para, ""]
        for bullet in section.bullets:
            out.append(f"- {bullet}")
        if section.bullets:
            out.append("")
        for table in section.tables:
            if table.title:
                out += [f"**{table.title}**", ""]
            out.append("| " + " | ".join(table.headers) + " |")
            out.append("|" + "|".join("---" for _ in table.headers) + "|")
            for row in table.rows:
                out.append("| " + " | ".join(str(c) for c in row) + " |")
            out.append("")
            if table.note:
                out += [f"*{table.note}*", ""]

    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    with open(path, "w") as f:
        f.write("\n".join(out).rstrip() + "\n")
    return path


def write_docx(res: Results, path: str) -> str:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    title, meta = _memo_header(res)
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    heading = doc.add_paragraph()
    run = heading.add_run(title)
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x18, 0x29, 0x57)

    for key, value in meta:
        p = doc.add_paragraph()
        r = p.add_run(f"{key}: ")
        r.bold = True
        p.add_run(value)
        p.paragraph_format.space_after = Pt(2)

    doc.add_paragraph("_" * 78)

    for section in build_sections(res):
        h = doc.add_heading(section.heading, level=1)
        for r in h.runs:
            r.font.color.rgb = RGBColor(0x18, 0x29, 0x57)
        for para in section.paragraphs:
            p = doc.add_paragraph(para)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for bullet in section.bullets:
            doc.add_paragraph(bullet, style="List Bullet")
        for table in section.tables:
            if table.title:
                p = doc.add_paragraph()
                p.add_run(table.title).bold = True
            t = doc.add_table(rows=1, cols=len(table.headers))
            t.style = "Light Grid Accent 1"
            for i, header in enumerate(table.headers):
                cell = t.rows[0].cells[i]
                cell.text = str(header)
                for para in cell.paragraphs:
                    for r in para.runs:
                        r.bold = True
                        r.font.size = Pt(8)
            for row in table.rows:
                cells = t.add_row().cells
                for i, value in enumerate(row):
                    cells[i].text = str(value)
                    for para in cells[i].paragraphs:
                        for r in para.runs:
                            r.font.size = Pt(8)
            if table.note:
                p = doc.add_paragraph()
                r = p.add_run(table.note)
                r.italic = True
                r.font.size = Pt(8)

    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    doc.save(path)
    return path


def build_memo(res: Results, base_path: str) -> list[str]:
    """Write memo.md and memo.docx next to each other; return both paths."""
    root = base_path[:-3] if base_path.endswith(".md") else base_path
    return [write_markdown(res, root + ".md"), write_docx(res, root + ".docx")]
